# -*- coding: utf-8 -*-
"""Maintainer tool: cut assets/kit.pptx out of the original 参考模板 template.

The kit keeps the template's master, layout and theme, plus three harvest slides:
  1. header  – title text box, hand icon, section tag, title rule, footer rule, disclaimer, signature
  2. icons   – the template's own icon graphics (svg + png fallback), one picture per key
  3. photos  – template photos, downscaled, one picture per key
Every harvested element is named "<kind>:<key>" so the engine can find it.

Usage:  python make_kit.py "参考模板" ../assets/kit.pptx
"""
import copy
import io
import sys

from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches

# (template slide no., shape id) of every icon the layouts may use
ICONS = {
    "battery": (56, 349), "home": (46, 1029), "growth": (11, 288), "coin": (3, 259),
    "ai": (10, 285), "city": (11, 289), "globe": (7, 264), "storm": (13, 227),
    "family": (67, 382), "man": (67, 370), "target": (39, 347), "check": (1, 1087),
    "box": (1, 1085), "cycle": (3, 400), "trend": (3, 394), "improve": (43, 34),
    "cart": (79, 79), "warn": (80, 218), "cube": (7, 185), "crown": (14, 47),
    "star": (23, 1058), "chat": (3, 399), "computer": (43, 33), "tag": (3, 257),
    "group": (77, 134), "women": (67, 417), "kid": (67, 413),
}
PHOTOS = {"car3d": (20, 31), "sedan": (22, 3074), "dealer": (22, 1042), "city_sun": (34, 3078),
          "city_dusk": (34, 3074), "city_tech": (35, 1051), "tower": (35, 1050)}
HEADER_SLIDE, RULE_SLIDE = 13, 10
HEADER_NAMES = {"文本框 2": "title", "图形 12": "icon", "文本框 1": "tag", "textbox 292": "sign"}
RULE_NAMES = {"直线连接符 4": "rule", "Grey bottom line": "foot_rule"}
PHOTO_MAX_PX = 1600


def find(shapes, sid):
    for sh in shapes:
        if sh.shape_type == 6:
            r = find(sh.shapes, sid)
            if r is not None:
                return r
        elif sh.shape_id == sid:
            return sh._element
    return None


def paste(dst_slide, el, src_part, name):
    el = copy.deepcopy(el)
    for node in el.iter():
        for attr in (qn("r:embed"), qn("r:link")):
            rid = node.get(attr)
            if rid:
                node.set(attr, dst_slide.part.relate_to(src_part.related_part(rid), RT.IMAGE))
    el.find(".//" + qn("p:cNvPr")).set("name", name)
    dst_slide.shapes._spTree.append(el)


def main(src, dst):
    prs = Presentation(src)
    slides = list(prs.slides)
    layout = prs.slide_layouts[0]
    hdr, ico, pho = (prs.slides.add_slide(layout) for _ in range(3))

    s = slides[HEADER_SLIDE - 1]
    for sh in s.shapes:
        if sh.name in HEADER_NAMES:
            paste(hdr, sh._element, s.part, "hdr:" + HEADER_NAMES[sh.name])
        elif sh.name == "Disclaimer" and sh.has_text_frame and "本报告" in sh.text_frame.text:
            paste(hdr, sh._element, s.part, "hdr:disc")
    s = slides[RULE_SLIDE - 1]
    for sh in s.shapes:
        if sh.name in RULE_NAMES:
            paste(hdr, sh._element, s.part, "hdr:" + RULE_NAMES[sh.name])

    for key, (sno, sid) in ICONS.items():
        sl = slides[sno - 1]
        el = find(sl.shapes, sid)
        if el is None:
            raise SystemExit(f"icon {key} not found on slide {sno}")
        paste(ico, el, sl.part, "icon:" + key)

    for i, (key, (sno, sid)) in enumerate(PHOTOS.items()):
        sl = slides[sno - 1]
        el = find(sl.shapes, sid)
        blob = sl.part.related_part(el.find(".//" + qn("a:blip")).get(qn("r:embed"))).blob
        im = Image.open(io.BytesIO(blob))
        im.thumbnail((PHOTO_MAX_PX, PHOTO_MAX_PX))
        buf = io.BytesIO()
        if im.mode in ("RGBA", "LA", "P"):
            im.convert("RGBA").save(buf, "PNG", optimize=True)
        else:
            im.convert("RGB").save(buf, "JPEG", quality=88)
        buf.seek(0)
        pic = pho.shapes.add_picture(buf, Inches(0.2 + i * 1.8), Inches(1), Inches(1.6))
        pic.name = "photo:" + key

    lst = prs.slides._sldIdLst
    for sid in list(lst)[:len(slides)]:
        prs.part.drop_rel(sid.rId)
        lst.remove(sid)
    prs.save(dst)
    print("kit written:", dst)


if __name__ == "__main__":
    main(*sys.argv[1:3])
