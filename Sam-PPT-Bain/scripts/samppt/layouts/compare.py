# -*- coding: utf-8 -*-
"""Comparison pages: versus, spec table, radar, 2x2 matrix, butterfly, persona."""
from lxml import etree
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE as DASH
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from ..charts import *
from ..components import *
from ..deck import grad, alpha, shadow, rgb, set_font
from ..tokens import *


def frame(D, s):
    D.new(s["title"], s.get("section", ""), notes=s.get("speaker_notes"))


# =================================================================================== versus
def versus(D, s):
    """Two sides (pros / cons / fit + photo) facing a VS badge, and a mirrored metric bar chart."""
    frame(D, s)
    gap = 0.53
    w = (CW - gap) / 2
    for i, sd in enumerate(s["sides"][:2]):
        x = EL if i == 0 else ER - w
        tone = sd.get("tone", "navy" if i == 0 else "red")
        col = col_of(tone)
        wcard(D, x, CARD_TOP, w, 2.9, tone)
        gbar(D, x, CARD_TOP, w, 0.45, [(sd["name"] + "  ", {"face": H}), ("｜" + sd.get("claim", ""), {})], tone, 13)
        if sd.get("photo"):
            if i == 0:
                D.photo_fit(sd["photo"], x + 0.12, 2.05, 1.95)
            else:
                D.photo_fit(sd["photo"], x + w - 2.0, 1.9, 1.9)
        lx = x + 2.1 if i == 0 else x + 0.25
        for k, (head, items, ic, icol) in enumerate((("优势", sd.get("pros", []), "check", col),
                                                       ("劣势", sd.get("cons", []), "warn", GREY_M))):
            yy = 1.78 + k * 0.92
            chip(D, lx, yy + 0.02, 0.62, head, tone if k == 0 else "navy", size=10)
            for j, it in enumerate(items[:2]):
                D.icon(ic, lx + 0.78, yy + 0.03 + j * 0.38, 0.24, color=icol)
                D.text(lx + 1.1, yy + j * 0.38, 2.75, 0.32, it, size=11, color=GREY, anchor="m")
        D.line(x + 0.2, 3.62, x + w - 0.2, 3.62, NAVY_B if tone == "navy" else RED_B, 0.5, DASH.DASH)
        if sd.get("fit"):
            D.text(x + 0.25, 3.66, w - 0.5, 0.32, sd["fit"], size=10.5, color=col, face=M, anchor="m", align="l" if i == 0 else "r")
    gcircle(D, (EL + ER) / 2 - 0.5, 2.1, 1.0, "red", text="VS", size=22, ring=3)
    m = s.get("metrics")
    if not m:
        return
    wcard(D, EL, 4.25, CW, 2.3)
    gbar(D, EL, 4.25, CW, 0.42, m["title"], "navy", 13)
    rows = m["rows"][:5]
    better = [((r["a"] < r["b"]) if r.get("better", "low") == "low" else (r["a"] > r["b"])) for r in rows]
    D.text(EL + 0.3, 4.72, 3.0, 0.26, [("◀ " + m.get("left", "") + "  ", {"face": M}), (f"{sum(better)} 项占优", {"face": B, "color": NAVY})],
           size=10.5, color=NAVY)
    D.text(ER - 3.3, 4.72, 3.0, 0.26, [(f"{len(rows) - sum(better)} 项占优", {"face": B, "color": RED}), ("  " + m.get("right", "") + " ▶", {"face": M})],
           size=10.5, color=RED, align="r")
    cx, half = (EL + ER) / 2, 3.9
    for j, r in enumerate(rows):
        y = 5.02 + j * 0.3
        mx = max(r["a"], r["b"])
        wa, wb = half * r["a"] / mx, half * r["b"] / mx
        D.text(cx - 1.0, y, 2.0, 0.24, r["label"], size=10, color=NAVY_D, face=M, align="c", anchor="m")
        grad(D.rect(cx - 1.05 - wa, y + 0.04, wa, 0.16, fill=NAVY), [(0, "2E5A94", 100), (100, "16365F", 100)], angle=0)
        grad(D.rect(cx + 1.05, y + 0.04, wb, 0.16, fill=RED), GRAD_RED, angle=0)
        D.text(cx - 1.05 - wa - 0.62, y, 0.57, 0.24, f'{r["a"]:g}', size=10, color=NAVY, face=B, align="r", anchor="m")
        D.text(cx + 1.05 + wb + 0.05, y, 0.6, 0.24, f'{r["b"]:g}', size=10, color=RED, face=B, anchor="m")
        if better[j]:
            D.box_text(cx - 1.05 - wa - 1.05, y + 0.02, 0.4, 0.2, "优", NAVY, WHITE, 8.5, B, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.5)
        else:
            D.box_text(cx + 1.05 + wb + 0.62, y + 0.02, 0.4, 0.2, "优", RED, WHITE, 8.5, B, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.5)


# =================================================================================== spec_table
def spec_table(D, s):
    """Native table with a framed focus column; hero statement, mid-value mini bars and one finding on the right."""
    frame(D, s)
    rows = [s["columns"]] + s["rows"]
    nr, nc = len(rows), len(s["columns"])
    tx, tw = EL, 8.68
    th = 5.38
    rh0 = 0.55
    rh = (th - rh0) / (nr - 1)
    plate = D.rect(tx, CARD_TOP, tw, th, fill=WHITE)
    shadow(plate, opacity=12, blur=6, dist=1.5)
    shp = D.s.shapes.add_table(nr, nc, Inches(tx), Inches(CARD_TOP), Inches(tw), Inches(th))
    tbl = shp.table
    c0 = 1.55
    widths = [c0] + [(tw - c0) / (nc - 1)] * (nc - 1)
    for c, w in enumerate(widths):
        tbl.columns[c].width = Inches(w)
    focus = s.get("focus_col")
    hdr_cols = [NAVY_D] + [NAVY, RED, NAVY_M, NAVY][: nc - 1]
    if focus:
        hdr_cols = [NAVY_D] + [RED if c == focus else (NAVY if c % 2 else NAVY_M) for c in range(1, nc)]
    hi = {tuple(x) for x in s.get("highlight_cells", [])}
    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(rh0 if r == 0 else rh)
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(hdr_cols[c] if r == 0 else (NAVY_T if c == 0 else ("FBF3F1" if c == focus else WHITE)))
            cell.margin_left = Inches(0.42 if (c == 0 and r) else 0.08); cell.margin_right = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame; tf.word_wrap = True
            p_ = tf.paragraphs[0]; p_.alignment = PP_ALIGN.LEFT if (c == 0 and r) else PP_ALIGN.CENTER
            run = p_.add_run(); run.text = val
            if r == 0:
                set_font(run, M, 13, WHITE)
            elif c == 0:
                set_font(run, M, 11.5, NAVY)
            else:
                set_font(run, B if (r, c) in hi else R, 11, RED if (r, c) in hi else GREY)
    tel = shp._element
    for tc in tel.iter(qn("a:tcPr")):
        for side in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
            ln = etree.SubElement(tc, qn(side)); ln.set("w", "6350")
            etree.SubElement(etree.SubElement(ln, qn("a:solidFill")), qn("a:srgbClr")).set("val", "D5DDE8")
        f_ = tc.find(qn("a:solidFill"))
        if f_ is not None:
            tc.remove(f_); tc.append(f_)
    tp = tel.find(".//" + qn("a:tblPr")); tp.set("firstRow", "0"); tp.set("bandRow", "0")
    st = tp.find(qn("a:tableStyleId"))
    if st is not None:
        tp.remove(st)
    for r, ic in enumerate(s.get("row_icons", []), start=1):
        if ic:
            D.icon(ic, tx + 0.12, CARD_TOP + rh0 + (r - 1) * rh + (rh - 0.26) / 2, 0.26, color=NAVY)
    if focus:
        fx = tx + sum(widths[:focus])
        fr = D.rect(fx - 0.03, CARD_TOP - 0.05, widths[focus] + 0.06, th + 0.1, line=RED, lw=2.25,
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.03)
        shadow(fr, color=RED, opacity=20)
        if s.get("focus_label"):
            chip(D, fx + (widths[focus] - 1.6) / 2, CARD_TOP + th - 0.11, 1.6, s["focus_label"], "red", h=0.3, size=10.5)
    rx, rw = ER - 3.58, 3.58
    hero = s.get("hero")
    if hero:
        hb = D.rect(rx, CARD_TOP, rw, 1.5, fill=RED)
        grad(hb, GRAD_RED, angle=0); shadow(hb, color=RED, opacity=28, blur=10, dist=3)
        c = D.rect(rx + 0.18, 1.42, 0.85, 0.85, fill=WHITE, shape=MSO_SHAPE.OVAL); shadow(c, color=RED, opacity=30)
        D.icon(hero.get("icon", "battery"), rx + 0.32, 1.56, 0.57, color=RED)
        D.text(rx + 1.2, 1.28, rw - 1.3, 0.3, hero.get("kicker", ""), size=11.5, color=ON_RED, face=M)
        D.text(rx + 1.2, 1.58, rw - 1.25, 0.45, hero["title"], size=18, color=WHITE, face=H)
        D.text(rx + 1.2, 2.1, rw - 1.3, 0.3, hero.get("sub", ""), size=10.5, color=ON_RED)
    mb = s.get("mini_bars")
    if mb:
        wcard(D, rx, 2.85, rw, 2.25)
        gbar(D, rx, 2.85, rw, 0.38, mb["title"], "navy", 11.5)
        cols_ = (NAVY, RED, NAVY_L)
        for g, grp in enumerate(mb["groups"][:2]):
            gy = 3.33 + g * 0.88
            D.text(rx + 0.15, gy, rw - 0.3, 0.24, grp["label"], size=9.5, color=GREY_M, face=M)
            mx = max(grp["values"])
            for k, v in enumerate(grp["values"][:3]):
                yy = gy + 0.26 + k * 0.19
                D.text(rx + 0.15, yy - 0.04, 0.45, 0.2, mb["names"][k], size=9, color=GREY)
                D.rect(rx + 0.65, yy, 2.2 * v / mx, 0.12, fill=cols_[k])
                D.text(rx + 0.65 + 2.2 * v / mx + 0.05, yy - 0.05, 0.6, 0.22, f"{v:g}", size=9, color=cols_[k] if k < 2 else NAVY_M, face=B)
    f = s.get("finding")
    if f:
        tab_card(D, rx, 5.28, rw, 1.27, "01", f["title"], f.get("metric"), f.get("tone", "navy"))
        bullets(D, rx + 0.18, 5.7, rw - 0.3, f.get("points", [])[:2], size=10.5, gap=0.34)


# =================================================================================== radar
def radar(D, s):
    """Filled radar of 2–4 options, per-dimension leader chips, average-score circles and findings."""
    frame(D, s)
    dims, series = s["dims"], s["series"]
    names = list(series)
    cols_ = (NAVY, RED, "7C97BF", NAVY_M)
    lw = 7.98
    wcard(D, EL, CARD_TOP, lw, 5.4)
    gbar(D, EL, CARD_TOP, lw, 0.45, s.get("chart_title", "能力评分"), "navy", 13)
    cd = CategoryChartData(); cd.categories = dims
    for k in names:
        cd.add_series(k, series[k])
    ch = D.chart(XL_CHART_TYPE.RADAR_FILLED, cd, EL + 0.25, 1.95, lw - 0.5, 3.75)
    chart_text(ch, 10)
    for i, ser in enumerate(ch.plots[0].series):
        c = cols_[i % 4]
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(c)
        ser_alpha(ser, 18)
        ser.format.line.color.rgb = rgb(c); ser.format.line.width = Pt(2.25)
    smax = s.get("max", 10)
    ch.value_axis.maximum_scale = smax; ch.value_axis.minimum_scale = 0; ch.value_axis.major_unit = smax / 5
    gridlines(ch.value_axis, "D5DDE8")
    ch.value_axis.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    legend(D, EL + 0.3, 1.7, names, cols_)
    D.line(EL + 0.2, 5.72, EL + lw - 0.2, 5.72, NAVY_B, 0.5, DASH.DASH)
    D.text(EL + 0.2, 5.8, 1.0, 0.6, ["各维度", "领先者"], size=9.5, color=GREY_M, face=M, spacing=1.05)
    sw = (lw - 1.25) / len(dims)
    for i, d in enumerate(dims):
        vals = [series[n][i] for n in names]
        best = [n for n in names if series[n][i] == max(vals)]
        x = EL + 1.05 + i * sw
        D.text(x, 5.8, sw - 0.07, 0.24, d, size=9, color=GREY, align="c")
        chip(D, x + 0.05, 6.07, sw - 0.17, "/".join(best), "navy" if best == [names[0]] else "red", h=0.26, size=9)
    rx, rw = 8.6, ER - 8.6
    wcard(D, rx, CARD_TOP, rw, 1.75)
    D.text(rx + 0.18, 1.22, 2.0, 0.26, "六项均分" if len(dims) == 6 else "平均得分", size=10, color=GREY_M, face=M)
    avgs = [sum(series[n]) / len(dims) for n in names]
    top = avgs.index(max(avgs))
    xs = [rx + 0.25 + k * (rw - 0.5) / len(names) for k in range(len(names))]
    for k, (n, a) in enumerate(zip(names, avgs)):
        big = k == top
        d = 1.12 if big else 0.92
        slot = (rw - 0.5) / len(names)
        x = xs[k] + (slot - d) / 2
        y = 1.42 if big else 1.6
        gcircle(D, x, y, d, "red" if big else "navy", text=f"{a:.1f}", size=19 if big else 16, ring=2)
        D.text(x - 0.2, y + d + 0.05, d + 0.4, 0.24, n, size=10, color=GREY, face=M, align="c")
        if big:
            chip(D, x + d - 0.25, 1.27, 0.9, s.get("top_label", "最均衡"), "red", h=0.26, size=9.5)
    findings(D, rx, 3.08, rw, 3.47, s.get("findings", []))


# =================================================================================== matrix
def matrix(D, s):
    """2x2 matrix with axes; one row framed as a highlighted band; hero statistic + two findings on the right."""
    frame(D, s)
    ox, oy, qw, qh, g = 1.25, 1.3, 3.6, 2.3, 0.12
    band = s.get("band")
    if band:
        by = oy + band.get("row", 1) * (qh + g)
        b = D.rect(ox - 0.12, by - 0.12, 2 * qw + g + 0.24, qh + 0.24, fill=RED_T, line=RED, lw=1.25,
                   shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.05, dash=DASH.DASH)
        alpha(b, 70)
    for k, q in enumerate(s["quads"][:4]):
        cx_, cy_ = k % 2, k // 2
        x, y = ox + cx_ * (qw + g), oy + cy_ * (qh + g)
        tone = q.get("tone", "navy")
        col = col_of(tone)
        qd = wcard(D, x, y, qw, qh, tone)
        grad(qd, [(0, "EAF0F7" if tone == "navy" else "FBEDEA", 100), (100, "FFFFFF", 100)], angle=45)
        gcircle(D, x + 0.25, y + 0.28, 0.62, tone, icon=q.get("icon", "target"))
        D.text(x + 1.02, y + 0.33, qw - 1.2, 0.5, q["title"], size=17, color=col, face=H, anchor="m")
        D.text(x + 0.27, y + 1.1, qw - 0.5, 0.6, q.get("desc", ""), size=12, color=GREY, spacing=1.2)
        D.line(x + 0.27, y + qh - 0.55, x + qw - 0.27, y + qh - 0.55, NAVY_B if tone == "navy" else RED_B, 0.5, DASH.DASH)
        D.text(x + 0.27, y + qh - 0.48, qw - 0.5, 0.3, q.get("label", ""), size=10, color=GREY_M, face=M)
    if band and band.get("label"):
        chip(D, ox + 2 * qw + g - 2.75, oy + 2 * qh + g, 2.75, band["label"], "red", h=0.3, size=9.5)
    ax = s.get("axes", {})
    D.line(ox - 0.3, oy + 2 * qh + g + 0.42, ox + 2 * qw + g + 0.3, oy + 2 * qh + g + 0.42, NAVY, 1.5, arrow=True)
    D.line(ox - 0.3, oy + 2 * qh + g + 0.42, ox - 0.3, oy - 0.1, NAVY, 1.5, arrow=True)
    D.text(ox, oy + 2 * qh + g + 0.48, 2 * qw + g, 0.28, ax.get("x", ""), size=11, color=NAVY, face=M, align="c")
    D.text(0.4, oy, 0.5, 2 * qh + g, ax.get("y", ""), size=11, color=NAVY, face=M, align="c", anchor="m", vert=True)
    rx, rw = ER - 3.88, 3.88
    h = s.get("hero")
    if h:
        hb = D.rect(rx, CARD_TOP, rw, 2.15, fill=RED)
        grad(hb, GRAD_RED, angle=0); shadow(hb, color=RED, opacity=28, blur=10, dist=3)
        if h.get("photo"):
            D.photo(h["photo"], rx, 1.9, rw, 1.4, opacity=16, feather=0.32)
        c = D.rect(rx + 0.25, 1.42, 0.9, 0.9, fill=WHITE, shape=MSO_SHAPE.OVAL); shadow(c, color=RED, opacity=30)
        D.icon(h.get("icon", "home"), rx + 0.4, 1.57, 0.6, color=RED)
        D.text(rx + 1.35, 1.25, 2.4, 0.85, [(h["value"], {}), (h.get("unit", ""), {"size": 20})], size=44, color=WHITE, face=H)
        D.text(rx + 0.25, 2.42, rw - 0.4, 0.32, h.get("title", ""), size=15, color=WHITE, face=M)
        D.text(rx + 0.25, 2.82, rw - 0.4, 0.3, h.get("sub", ""), size=11.5, color=ON_RED, face=M)
    findings(D, rx, 3.5, rw, 3.05, s.get("findings", [])[:2], top_pad=0.55)


# =================================================================================== butterfly
def butterfly(D, s):
    """Mirrored bar chart of two groups across factors, difference pills, findings column."""
    frame(D, s)
    lw = 8.33
    wcard(D, EL, CARD_TOP, lw, 5.4)
    gbar(D, EL, CARD_TOP, lw, 0.45, s.get("chart_title", ""), "navy", 13)
    A, Bg = s["left"], s["right"]
    gcircle(D, EL + 0.25, 1.72, 0.42, "navy", icon=A.get("icon", "man"), ring=1.5)
    D.text(EL + 0.75, 1.72, 1.5, 0.42, A["name"], size=11, color=NAVY, face=M, anchor="m")
    gcircle(D, EL + 5.83, 1.72, 0.42, "red", icon=Bg.get("icon", "family"), ring=1.5)
    D.text(EL + 4.15, 1.72, 1.6, 0.42, Bg["name"], size=11, color=RED, face=M, anchor="m", align="r")
    D.text(EL + 6.97, 1.72, 1.1, 0.42, "差值", size=10, color=GREY_M, face=M, anchor="m", align="c")
    facs, a, b = s["factors"], A["values"], Bg["values"]
    cd = CategoryChartData(); cd.categories = facs
    cd.add_series(A["name"], [-v for v in a]); cd.add_series(Bg["name"], b)
    cx_, cy_, cw_, chh = EL + 0.1, 2.2, 6.85, 4.25
    lx, ty, lw_, th = 0.17, 0.02, 0.81, 0.96
    ch = D.chart(XL_CHART_TYPE.BAR_CLUSTERED, cd, cx_, cy_, cw_, chh)
    chart_text(ch, 10)
    plot_layout(ch, lx, ty, lw_, th)
    pl_ = ch.plots[0]; pl_.overlap = 100; pl_.gap_width = 55
    for i, ser in enumerate(pl_.series):
        ser.invert_if_negative = False
        ser_grad(ser, *(("16365F", "2E5A94") if i == 0 else ("C0503E", "8E2B1D")), angle=0)
        ser.data_labels.show_value = True
        ser.data_labels.number_format = "0;0"; ser.data_labels.number_format_is_linked = False
        ser.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        label_font(ser.data_labels, 10, (NAVY, RED)[i], face=B)
    ca = ch.category_axis
    ca.reverse_order = True; ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    axis_line(ca)
    vm = max(a + b) * 1.25
    ch.value_axis.minimum_scale = -vm; ch.value_axis.maximum_scale = vm
    hide_val_axis(ch)
    pt_, ph_ = cy_ + ty * chh, th * chh
    hot = s.get("diff_hot", 15)
    for i, (x, y) in enumerate(zip(a, b)):
        d = y - x
        yy = pt_ + (i + 0.5) * ph_ / len(facs)
        outline_pill(D, EL + 7.1, yy - 0.15, 0.85, f"{d:+g}", "red" if d > 0 else "navy", h=0.3, size=10, filled=abs(d) >= hot)
    findings(D, ER - 3.93, CARD_TOP, 3.93, 5.4, s.get("findings", []), top_pad=0.52)


# =================================================================================== persona
def persona(D, s):
    """Persona panel (photo, avatar, tags, quote) | attributes + needs | index bars against a baseline."""
    frame(D, s)
    pw = 3.35
    pnl = D.rect(EL, CARD_TOP, pw, 5.4, fill=NAVY)
    grad(pnl, [(0, "16365F", 100), (100, "2E5A94", 100)], angle=90); shadow(pnl, opacity=30, blur=10, dist=3)
    if s.get("photo"):
        D.photo(s["photo"], EL, CARD_TOP, pw, 1.55, opacity=40, feather=0.08)
    av = D.rect(EL + (pw - 1.5) / 2, 2.05, 1.5, 1.5, fill=WHITE, line=WHITE, lw=3, shape=MSO_SHAPE.OVAL)
    shadow(av, opacity=40, blur=12, dist=4)
    D.icon(s.get("avatar_icon", "family"), EL + (pw - 1.0) / 2, 2.3, 1.0, color=RED)
    D.text(EL, 3.65, pw, 0.5, s["name"], size=24, color=WHITE, face=H, align="c")
    D.text(EL + 0.1, 4.15, pw - 0.2, 0.3, s.get("tagline", ""), size=10, color=ON_DARK, align="c")
    for i, t in enumerate(s.get("tags", [])[:4]):
        alpha(D.box_text(EL + 0.27 + (i % 2) * 1.45, 4.58 + (i // 2) * 0.42, 1.3, 0.32, t, WHITE, WHITE, 10, M,
                         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.5), 16)
    if s.get("quote"):
        q = D.rect(EL + 0.2, 5.5, pw - 0.4, 0.88, fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT, adj=[-0.2, -0.75, 0.16667])
        shadow(q, opacity=30)
        D.text(EL + 0.35, 5.52, 0.4, 0.4, "“", size=26, color=RED, face=H)
        D.text(EL + 0.67, 5.55, pw - 0.95, 0.8, s["quote"], size=10.5, color=NAVY, face=M, anchor="m", spacing=1.15)
    mx, mw = EL + pw + 0.2, 4.4
    wcard(D, mx, CARD_TOP, mw, 2.4)
    gbar(D, mx, CARD_TOP, mw, 0.42, s.get("attrs_title", "基础属性"), "navy", 13)
    for i, at in enumerate(s.get("attrs", [])[:4]):
        xx, yy = mx + 0.18 + (i % 2) * 2.12, 1.75 + (i // 2) * 0.85
        gcircle(D, xx, yy + 0.06, 0.52, "navy", icon=at.get("icon", "tag"), ring=1.5)
        D.text(xx + 0.62, yy, 1.45, 0.26, at["label"], size=9.5, color=GREY_M)
        D.text(xx + 0.62, yy + 0.25, 1.5, 0.4, at["value"], size=13.5, color=NAVY, face=B)
    wcard(D, mx, 3.72, mw, 2.83, "red")
    gbar(D, mx, 3.72, mw, 0.42, s.get("needs_title", "核心诉求"), "red", 13)
    for i, nd in enumerate(s.get("needs", [])[:4]):
        yy = 4.27 + i * 0.56
        gcircle(D, mx + 0.18, yy, 0.44, "red", icon=nd.get("icon", "target"), ring=1.25)
        D.text(mx + 0.75, yy, 2.55, 0.44, nd["text"], size=11, color=GREY, anchor="m")
        if nd.get("tag"):
            D.box_text(mx + 3.3, yy + 0.09, 0.95, 0.26, nd["tag"], RED_T, RED, 9, M, line=RED_B, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.5)
    ix = mx + mw + 0.2
    iw = ER - ix
    t = s["index"]
    wcard(D, ix, CARD_TOP, iw, 5.4)
    gbar(D, ix, CARD_TOP, iw, 0.42, t["title"], "navy", 12.5)
    base = t.get("baseline", 100)
    vals = t["values"]
    cd = CategoryChartData(); cd.categories = t["categories"]; cd.add_series("index", vals)
    cx_, cy_, cw_, chh = ix + 0.1, 1.75, iw - 0.18, 4.15
    lx, ty, lw_, th = 0.36, 0.02, 0.58, 0.96
    ch = D.chart(XL_CHART_TYPE.BAR_CLUSTERED, cd, cx_, cy_, cw_, chh)
    chart_text(ch, 10)
    plot_layout(ch, lx, ty, lw_, th)
    ser = ch.plots[0].series[0]
    ser_grad(ser, "C0503E", "8E2B1D", angle=0)
    ch.plots[0].gap_width = 70
    ch.category_axis.reverse_order = True; axis_line(ch.category_axis)
    vmax = base * 2
    ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = vmax
    hide_val_axis(ch)
    ser.data_labels.show_value = True; ser.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    label_font(ser.data_labels, 11, NAVY, face=B)
    for j, v in enumerate(vals):
        if v < base:
            ser.points[j].format.fill.solid(); ser.points[j].format.fill.fore_color.rgb = rgb(NAVY_L)
            pdl = ser.points[j].data_label; pdl.position = XL_LABEL_POSITION.INSIDE_END; label_font(pdl, 11, NAVY, face=B)
    xb = cx_ + (lx + lw_ * base / vmax) * cw_
    D.line(xb, cy_ + 0.05, xb, cy_ + chh - 0.05, RED, 1, DASH.DASH)
    D.box_text(xb - 0.5, cy_ + chh + 0.02, 1.0, 0.24, f"平均 = {base:g}", None, RED, 9, M)
    if t.get("note"):
        D.text(ix + 0.2, 6.2, iw - 0.4, 0.3, t["note"], size=10, color=GREY, align="c", hl={"face": M, "color": NAVY})


LAYOUTS = {"versus": versus, "spec_table": spec_table, "radar": radar, "matrix": matrix,
           "butterfly": butterfly, "persona": persona}
