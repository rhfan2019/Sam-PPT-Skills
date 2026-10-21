# -*- coding: utf-8 -*-
"""Bubble map, smile curve, pyramid, timeline, company cards, strategy house."""
from lxml import etree
from pptx.chart.data import BubbleChartData, CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_LINE_DASH_STYLE as DASH
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Pt

from ..charts import *
from ..components import *
from ..deck import grad, alpha, shadow, rgb, set_font
from ..tokens import *

LP = {"right": XL_LABEL_POSITION.RIGHT, "left": XL_LABEL_POSITION.LEFT, "above": XL_LABEL_POSITION.ABOVE,
      "below": XL_LABEL_POSITION.BELOW, "center": XL_LABEL_POSITION.CENTER}


def frame(D, s):
    D.new(s["title"], s.get("section", ""), notes=s.get("speaker_notes"))


# =================================================================================== bubble
def bubble(D, s):
    """Bubble map (x, y, size) with shaded value zones, a call-out on the leader, findings column."""
    frame(D, s)
    lw = 8.6
    wcard(D, EL, CARD_TOP, lw, 5.4)
    gbar(D, EL, CARD_TOP, lw, 0.45, s.get("chart_title", ""), "navy", 13)
    ax = s.get("axes", {})
    xmax, ymax = ax.get("xmax", 35), ax.get("ymax", 5)
    cols_ = (NAVY, RED, "7C97BF", NAVY_M)
    cx_, cy_, cw_, chh = EL + 0.1, 2.0, lw - 0.2, 4.4
    lx, ty, lw_, th = 0.08, 0.02, 0.89, 0.86
    pl, pw, pt, ph = cx_ + lx * cw_, lw_ * cw_, cy_ + ty * chh, th * chh
    X = lambda v: pl + pw * v / xmax
    Y = lambda v: pt + ph * (1 - v / ymax)
    for z in s.get("zones", []):
        D.rect(X(z["from"]), pt, X(z["to"]) - X(z["from"]), ph, fill="EAF0F7" if z.get("tone", "navy") == "navy" else "FBEDEA")
        chip(D, X(z["from"]) + 0.05, pt + 0.05, X(z["to"]) - X(z["from"]) - 0.1, z["label"], z.get("tone", "navy"), h=0.26, size=9)
    series = s["series"]
    cd = BubbleChartData()
    for k, pts in series.items():
        se = cd.add_series(k)
        for p in pts:
            se.add_data_point(p["x"], p["y"], p["size"])
    ch = D.chart(XL_CHART_TYPE.BUBBLE, cd, cx_, cy_, cw_, chh)
    chart_text(ch, 9.5)
    plot_layout(ch, lx, ty, lw_, th)
    bubble_scale(ch, s.get("bubble_scale", 32))
    for i, (k, pts) in enumerate(series.items()):
        ser = ch.plots[0].series[i]
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(cols_[i % 4])
        ser_alpha(ser, 88)
        ser.format.line.color.rgb = rgb(WHITE); ser.format.line.width = Pt(1.5)
        for j, p in enumerate(pts):
            dl = ser.points[j].data_label
            dl.has_text_frame = True; dl.text_frame.text = p["label"]; dl.position = LP[p.get("pos", "right")]
            set_font(dl.text_frame.paragraphs[0].runs[0], M, 9.5, NAVY_D)
    va, ca = ch.value_axis, ch.category_axis
    va.minimum_scale = 0; va.maximum_scale = ymax; va.major_unit = ax.get("ystep", 1)
    ca.minimum_scale = 0; ca.maximum_scale = xmax; ca.major_unit = ax.get("xstep", 5)
    gridlines(va); axis_line(va, show=False); axis_line(ca)
    legend(D, EL + 0.3, 1.68, list(series), cols_, round_=True)
    D.text(EL + 4.55, 1.68, 3.9, 0.26, f'横轴：{ax.get("x", "")}  纵轴：{ax.get("y", "")}', size=9, color=GREY_M, align="r", anchor="m")
    co = s.get("callout")
    if co:
        p = next(p for p in series[co["series"]] if p["label"] == co["point"])
        tx, ty_ = X(p["x"]) + 0.02, Y(p["y"]) + 0.05
        callout(D, min(tx + 0.4, EL + lw - 1.7), pt + 1.0, 1.55, 0.44, tx, ty_, co["text"], tone="navy", size=10)
    findings(D, ER - 3.58, CARD_TOP, 3.58, 5.4, s.get("findings", []), top_pad=0.55)


# =================================================================================== smile_curve
def smile_curve(D, s):
    """Value-added index along a chain (smoothed line), zones behind it, one card per link below."""
    frame(D, s)
    wcard(D, EL, CARD_TOP, CW, 5.4)
    links, vals = s["links"], s["values"]
    n = len(links)
    hi = set(s.get("highlight", [0, n - 1]))
    px, pw = EL + 0.3, CW - 0.6
    cw = pw / n
    cy_, chh = 1.3, 3.3
    ty, th = 0.14, 0.74
    pt, ph = cy_ + ty * chh, th * chh
    for z in s.get("zones", []):
        a, b = z["from"], z["to"] + 1
        tone = z.get("tone", "navy")
        D.rect(px + a * cw + 0.04, pt - 0.08, (b - a) * cw - 0.08, ph + 0.08, fill="FBEDEA" if tone == "red" else "F1F4F9")
        chip(D, px + a * cw + ((b - a) * cw - 1.5) / 2, 1.3, 1.5, z["label"], tone, h=0.28, size=10)
    cd = CategoryChartData(); cd.categories = links; cd.add_series(s.get("series_name", "附加值指数"), vals)
    ch = D.chart(XL_CHART_TYPE.LINE_MARKERS, cd, px, cy_, pw, chh)
    chart_text(ch, 10)
    plot_layout(ch, 0, ty, 1, th)
    ser = ch.plots[0].series[0]
    ser.smooth = True
    fill_series(ser, NAVY, line=True, width=3.5)
    ser.marker.style = XL_MARKER_STYLE.CIRCLE; ser.marker.size = 11
    ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(WHITE)
    ser.marker.format.line.color.rgb = rgb(NAVY); ser.marker.format.line.width = Pt(2.25)
    ser.data_labels.show_value = True; ser.data_labels.position = XL_LABEL_POSITION.ABOVE
    label_font(ser.data_labels, 12, NAVY, face=B)
    for j in hi:
        m = ser.points[j].marker
        m.style = XL_MARKER_STYLE.CIRCLE; m.size = 16
        m.format.fill.solid(); m.format.fill.fore_color.rgb = rgb(RED); m.format.line.color.rgb = rgb(WHITE)
        dlp = ser.points[j].data_label
        dlp.position = XL_LABEL_POSITION.ABOVE
        label_font(dlp, 15, RED, face=H)
    vmax = max(vals) * 1.22
    ch.value_axis.minimum_scale = 0; ch.value_axis.maximum_scale = vmax
    hide_val_axis(ch)
    ch.category_axis.visible = False
    co = s.get("callout")
    if co:
        i = co["at"]
        yv = pt + ph * (1 - vals[i] / vmax)
        callout(D, px + i * cw - 0.15, yv + 0.55, 1.95, 0.42, px + (i + 0.5) * cw, yv + 0.12, co["text"], tone="navy", size=10)
    if s.get("axis_note"):
        D.text(px + 0.1, pt + ph - 0.3, 2.5, 0.25, s["axis_note"], size=9, color=GREY_M)
    notes, icons = s.get("notes", [""] * n), s.get("icons", ["cube"] * n)
    for j in range(n):
        x = px + j * cw + 0.06
        tone = "red" if j in hi else "navy"
        wcard(D, x, 4.75, cw - 0.12, 1.6, tone)
        gbar(D, x, 4.75, cw - 0.12, 0.4, links[j], tone, 11.5)
        gcircle(D, x + (cw - 0.12) / 2 - 0.25, 5.22, 0.5, tone, icon=icons[j], ring=1.5)
        D.text(x + 0.05, 5.8, cw - 0.22, 0.45, notes[j], size=10, color=RED if j in hi else GREY, face=M if j in hi else R, align="c", anchor="m")


# =================================================================================== pyramid
def pyramid(D, s):
    """Layered pyramid (top = hardest), connector to a numbered description row per layer, rising axis."""
    frame(D, s)
    layers = s["layers"][:6]
    n = len(layers)
    blues = [[(0, "16365F", 100), (100, "2E5A94", 100)], [(0, "2E5A94", 100), (100, "4A6A99", 100)],
             [(0, "5D7DAA", 100), (100, "7C97BF", 100)], [(0, "8FA8CB", 100), (100, "B7C8DF", 100)],
             [(0, "A9BEDC", 100), (100, "D3DEEE", 100)]]
    pcx, ptop, pH, pW = 3.95, 1.55, 4.8, 5.5
    lh = pH / n
    for i, L in enumerate(layers):
        y = ptop + i * lh
        wt, wb = pW * i / n, pW * (i + 1) / n
        g = 0.06
        tone = "red" if i == 0 else "navy"
        stops = GRAD_RED if i == 0 else blues[min(i - 1, 4)]
        if i == 0:
            sh = D.rect(pcx - wb / 2, y, wb, lh - g, fill=RED, shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
            grad(sh, stops, angle=90); shadow(sh, color=RED, opacity=30, blur=8, dist=3)
            D.text(pcx - wb / 2, y + lh * 0.56, wb, 0.3, L["name"], size=10.5, color=WHITE, face=M, align="c", anchor="m")
        else:
            inset = (wb - wt) / 2
            sh = D.rect(pcx - wb / 2, y, wb, lh - g, fill=NAVY, shape=MSO_SHAPE.TRAPEZOID, adj=inset / min(wb, lh - g))
            grad(sh, stops, angle=90); shadow(sh, opacity=22, blur=8, dist=3)
            D.text(pcx - wb / 2, y, wb, lh - g, L["name"], size=14, color=WHITE if i < n - 1 else NAVY_D, face=M, align="c", anchor="m")
        edge = pcx + (wt + wb) / 4 + 0.12
        D.line(edge, y + lh / 2 - 0.03, 7.2, y + lh / 2 - 0.03, RED if i == 0 else NAVY_B, 0.75, DASH.DASH)
        wcard(D, 7.45, y + 0.06, ER - 7.45, lh - 0.18, tone)
        gcircle(D, 7.15, y + lh / 2 - 0.3, 0.6, tone, text=str(n - i), size=15)
        D.text(7.95, y + 0.16, 3.6, 0.32, L["name"], size=13.5, color=col_of(tone), face=H)
        D.text(7.95, y + 0.5, ER - 8.6, 0.32, L.get("desc", ""), size=11, color=GREY)
        if L.get("icon"):
            D.icon(L["icon"], ER - 0.66, y + lh / 2 - 0.27, 0.42, color=RED if i == 0 else NAVY_L)
    if s.get("crown", True):
        D.icon("crown", pcx - 0.25, ptop - 0.5, 0.5, color=RED)
    if s.get("axis"):
        arr = D.rect(0.55, 1.6, 0.42, 4.7, fill=RED, shape=MSO_SHAPE.UP_ARROW, adj=[0.45, 0.6])
        grad(arr, [(0, "A23424", 100), (100, "F3D5CF", 100)], angle=90)
        D.text(1.05, 2.5, 0.4, 2.8, s["axis"], size=12, color=RED, face=M, anchor="m", vert=True)


# =================================================================================== timeline
def timeline(D, s):
    """Era band, chevron years, icon badge + card per event (focus / forecast marks), conclusion bar."""
    frame(D, s)
    ev = s["events"][:7]
    n = len(ev)
    cw = CW / n
    eras = s.get("eras", [])
    if eras:
        x0 = EL
        for k, e in enumerate(eras):
            x1 = EL + cw * e.get("until", n) if k < len(eras) - 1 else ER
            red = k > 0
            b = D.rect(x0, 1.2, x1 - x0 + (0.08 if k < len(eras) - 1 else 0), 0.42, fill=RED_T if red else NAVY_T,
                       line=RED_B if red else NAVY_B, lw=0.5, shape=MSO_SHAPE.CHEVRON if k else MSO_SHAPE.PENTAGON, adj=0.25)
            grad(b, [(0, "F3D5CF", 100), (100, "FBF0EE", 100)] if red else [(0, "D3DEEE", 100), (100, "F5F8FC", 100)], angle=0)
            D.text(x0 + 0.2, 1.2, x1 - x0 - 0.4, 0.42, [(e["label"] + "  ", {"face": H}), (e.get("sub", ""), {})],
                   size=12, color=RED if red else NAVY, face=M, align="c", anchor="m")
            x0 = x1
    for i, e in enumerate(ev):
        x = EL + i * cw
        mark = e.get("mark")
        hl, fc = mark == "focus", mark == "forecast"
        tone = "red" if hl else "navy"
        last = i == n - 1
        sh = D.rect(x, 1.85, cw + (0 if last else 0.08), 0.6, fill=NAVY,
                    shape=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON, adj=0.35)
        grad(sh, GRAD_RED if hl else ([(0, "B7C8DF", 100), (100, "D3DEEE", 100)] if fc else GRAD_NAVY), angle=0)
        shadow(sh, color=RED if hl else NAVY, opacity=20)
        D.text(x + 0.15, 1.85, cw - 0.25, 0.6, e["year"], size=20, color=NAVY if fc else WHITE, face=H, align="c", anchor="m")
        D.line(x + cw / 2, 2.45, x + cw / 2, 2.72, RED if hl else NAVY_B, 1)
        card = D.rect(x + 0.08, 3.08, cw - 0.16, 1.85, fill=RED_T if hl else WHITE, line=RED if hl else NAVY_B,
                      lw=1.25 if hl else 0.75, dash=DASH.DASH if fc else None)
        shadow(card, color=RED if hl else NAVY, opacity=12, blur=6, dist=1.5)
        gcircle(D, x + cw / 2 - 0.3, 2.72, 0.6, tone, icon=e.get("icon", "star"), ring=2)
        lines = e["text"] if isinstance(e["text"], list) else e["text"].split("，")
        D.text(x + 0.18, 3.5, cw - 0.36, 1.1, lines, size=11.5, color=GREY, spacing=1.2, after=4)
        if hl:
            chip(D, x + 0.22, 4.5, 0.9, e.get("mark_label", "重点"), "red", h=0.28, size=10)
        if fc:
            D.box_text(x + 0.22, 4.5, 0.9, 0.28, e.get("mark_label", "预测"), WHITE, NAVY_M, 10, M, line=NAVY_M,
                       shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.5)
            if s.get("photo"):
                D.photo_fit(s["photo"], x + cw - 1.3, 4.22, 1.15)
    if s.get("summary"):
        summary(D, EL, 5.3, CW, 1.05, s["summary"], label=s.get("summary_label", "转折点"))


# =================================================================================== company_cards
def company_cards(D, s):
    """One card per player: route pill, big number, scale bar vs. the largest, three strengths; summary bar."""
    frame(D, s)
    cs = s["companies"][:6]
    mx = max(c["value"] for c in cs)
    for i, ((x, w), c) in enumerate(zip(cols(len(cs), 0.18), cs)):
        tone = c.get("tone", "navy" if i % 2 == 0 else "red")
        col = col_of(tone)
        wcard(D, x, CARD_TOP, w, 4.6, tone)
        gbar(D, x, CARD_TOP, w, 0.62, "", tone)
        D.text(x, 1.17, w, 0.36, c["name"], size=16, color=WHITE, face=H, align="c", anchor="m")
        gcircle(D, x + w / 2 - 0.27, 1.54, 0.54, tone, icon=c.get("icon", "cube"), ring=2.5)
        D.box_text(x + 0.25, 2.25, w - 0.5, 0.34, c.get("route", ""), WHITE, col, 11.5, M, line=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.5)
        D.text(x, 2.7, w, 0.75, [(f'{c["value"]:g}', {}), (" " + c.get("unit", ""), {"size": 12, "face": M, "color": GREY})],
               size=34, color=col, face=H, align="c")
        D.text(x, 3.42, w, 0.26, c.get("value_note", ""), size=10, color=GREY_M, align="c")
        progress(D, x + 0.3, 3.8, w - 0.6, c["value"] / mx, col, h=0.09)
        D.text(x + 0.3, 3.92, w - 0.6, 0.24, s.get("bench_text", "规模最大 = 100%") if c["value"] == mx
               else f'约为{s.get("bench_name", "第一")} {c["value"] / mx:.0%}', size=9, color=GREY_M, align="c")
        D.line(x + 0.25, 4.3, x + w - 0.25, 4.3, NAVY_B if tone == "navy" else RED_B, 0.75, DASH.DASH)
        for j, p in enumerate(c.get("points", [])[:3]):
            yy = 4.42 + j * 0.42
            D.rect(x + 0.3, yy + 0.11, 0.08, 0.08, fill=col)
            D.text(x + 0.48, yy, w - 0.6, 0.3, p, size=11.5, color=GREY, anchor="m")
    if s.get("summary"):
        summary(D, EL, 5.95, CW, 0.6, s["summary"], label=s.get("summary_label", "总结"))


# =================================================================================== strategy_house
def strategy_house(D, s):
    """Roof (vision + two KPI circles), 3–5 pillars with icon, points and tag, foundation bar."""
    frame(D, s)
    if s.get("photo"):
        D.photo(s["photo"], EL, 3.2, CW, 2.7, opacity=28, feather=0.18)
    roof = s["roof"]
    r = D.rect(EL, 1.1, CW, 1.55, fill=NAVY, shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
    grad(r, GRAD_NAVY, angle=90); shadow(r, opacity=30, blur=10, dist=3)
    mid = (EL + ER) / 2
    D.text(mid - 2.1, 1.85, 4.2, 0.3, roof.get("kicker", ""), size=11, color=NAVY_L, face=M, align="c")
    D.text(mid - 2.3, 2.12, 4.6, 0.4, roof["title"], size=15, color=WHITE, face=H, align="c")
    for cx, (v, lab) in zip((mid - 3.1, mid + 3.1), roof.get("kpis", [])[:2]):
        c = D.rect(cx - 0.6, 1.62, 1.2, 1.2, fill=WHITE, line=RED, lw=2, shape=MSO_SHAPE.OVAL)
        shadow(c, color=RED, opacity=30, blur=10, dist=3)
        D.text(cx - 0.6, 1.8, 1.2, 0.5, v, size=22, color=RED, face=H, align="c")
        D.text(cx - 0.6, 2.28, 1.2, 0.3, lab, size=9, color=NAVY, face=M, align="c")
    ps = s["pillars"][:5]
    for i, ((x, w), p) in enumerate(zip(cols(len(ps), 0.2, EL + 0.2, ER - 0.2), ps)):
        tone = "navy" if i % 2 == 0 else "red"
        wcard(D, x, 2.9, w, 2.75, tone)
        gbar(D, x, 2.9, w, 0.62, "", tone)
        D.text(x, 2.92, w, 0.36, f'{s.get("pillar_word", "支柱")} {i + 1}｜{p["name"]}', size=13.5, color=WHITE, face=M, align="c", anchor="m")
        gcircle(D, x + w / 2 - 0.3, 3.3, 0.6, tone, icon=p.get("icon", "cube"), ring=2.5)
        D.text(x + 0.2, 4.05, w - 0.4, 1.0, p.get("points", []), size=12, color=GREY, align="c", spacing=1.2, after=4)
        if p.get("tag"):
            D.line(x + 0.3, 5.08, x + w - 0.3, 5.08, NAVY_B if tone == "navy" else RED_B, 0.5, DASH.DASH)
            D.box_text(x + w / 2 - 0.75, 5.18, 1.5, 0.3, p["tag"], RED_T if tone == "red" else NAVY_T, col_of(tone), 10, M,
                       shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.5)
    D.rect(EL, 5.8, CW, 0.08, fill="BFBFBF")
    gbar(D, EL, 5.95, CW, 0.6, s.get("foundation", ""), "navy", 15)


LAYOUTS = {"bubble": bubble, "smile_curve": smile_curve, "pyramid": pyramid, "timeline": timeline,
           "company_cards": company_cards, "strategy_house": strategy_house}
