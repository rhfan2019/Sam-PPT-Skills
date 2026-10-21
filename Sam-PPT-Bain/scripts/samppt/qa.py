# -*- coding: utf-8 -*-
"""Automatic checks run after every build. Each returns a list of findings (dicts)."""
import math
import re
import zipfile

from pptx import Presentation
from pptx.util import Pt

from .tokens import EL, ER, TOP, SLIDE_H

EMU = 914400
FRAME_PREFIX = "frame:"
MIN_PT = 7.5


def _box(sh):
    return sh.left / EMU, sh.top / EMU, (sh.left + sh.width) / EMU, (sh.top + sh.height) / EMU


def _text_runs(sh):
    if not sh.has_text_frame:
        return []
    return [r for p in sh.text_frame.paragraphs for r in p.runs if r.text.strip()]


def check_edges(prs, tol=0.006):
    """Nothing may cross the content edges (the template title rule / footer edges)."""
    out = []
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.name.startswith(FRAME_PREFIX):
                continue
            x0, y0, x1, y1 = _box(sh)
            if x0 < EL - tol or x1 > ER + tol:
                out.append({"slide": i, "check": "edge", "shape": sh.name,
                            "detail": f"x {x0:.2f}–{x1:.2f} outside {EL}–{ER}"})
            if y1 > 7.16 - tol and not sh.name.startswith(FRAME_PREFIX):
                out.append({"slide": i, "check": "edge", "shape": sh.name, "detail": f"bottom {y1:.2f} below footer rule 7.16"})
    return out


def _char_w(ch, size):
    """Approximate glyph advance in inches for OPPOSans."""
    if "⺀" <= ch <= "鿿" or "＀" <= ch <= "￯" or "　" <= ch <= "〿":
        return size / 72 * 1.0
    if ch == " ":
        return size / 72 * 0.28
    return size / 72 * 0.56


def check_text_fit(prs, slack=1.12):
    """Estimate wrapped height of every text box and flag ones whose text is taller/wider than the box."""
    out = []
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.name.startswith(FRAME_PREFIX) or not sh.has_text_frame or not _text_runs(sh):
                continue
            tf = sh.text_frame
            w = sh.width / EMU
            h = sh.height / EMU
            if w < 0.05:
                continue
            wrap = tf.word_wrap is not False
            vert = tf._txBody.find("{http://schemas.openxmlformats.org/drawingml/2006/main}bodyPr").get("vert")
            if vert:
                w, h = h, w
            tot_h, max_line = 0.0, 0.0
            for p in tf.paragraphs:
                size = max([(r.font.size.pt if r.font.size else 11) for r in p.runs] or [11])
                width = sum(_char_w(c, (r.font.size.pt if r.font.size else 11)) for r in p.runs for c in r.text)
                lines = max(1, math.ceil(width / w - 0.02)) if wrap else 1
                max_line = max(max_line, width)
                ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
                tot_h += lines * size / 72 * 1.2 * ls
            name = sh.name
            if not wrap and max_line > w * slack:
                out.append({"slide": i, "check": "text-width", "shape": name, "detail": f"needs {max_line:.2f}in, box {w:.2f}in"})
            if tot_h > h * slack + 0.06:
                out.append({"slide": i, "check": "text-height", "shape": name,
                            "detail": f"~{tot_h:.2f}in of text in a {h:.2f}in box: “{tf.text[:24]}”"})
    return out


def check_min_font(prs):
    out = []
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.name.startswith(FRAME_PREFIX):
                continue
            for r in _text_runs(sh):
                if r.font.size and r.font.size.pt < MIN_PT:
                    out.append({"slide": i, "check": "min-font", "shape": sh.name, "detail": f"{r.font.size.pt}pt “{r.text[:12]}”"})
                    break
    return out


def check_xml(path):
    """Defects that pass schema checks but make PowerPoint refuse the file."""
    out = []
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n.startswith("ppt/charts/chart") and n.endswith(".xml"):
                x = z.read(n).decode("utf8")
                if re.search(r"<a:ea [^>]*/>\s*<a:ea ", x):
                    out.append({"slide": 0, "check": "xml", "shape": n, "detail": "duplicate <a:ea> in chart text"})
    return out


def run_all(path):
    prs = Presentation(path)
    res = check_edges(prs) + check_text_fit(prs) + check_min_font(prs) + check_xml(path)
    return prs, res
