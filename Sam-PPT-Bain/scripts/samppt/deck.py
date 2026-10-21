# -*- coding: utf-8 -*-
"""Deck: the template frame (master, header, footer) plus drawing primitives and finishing effects."""
import copy
import io
import os
import re

from lxml import etree
from PIL import Image, ImageChops, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE as DASH
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from .tokens import *

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
KIT = os.path.join(os.path.dirname(__file__), "..", "..", "assets", "kit.pptx")
EMU = 914400


def rgb(h):
    return RGBColor.from_string(h)


def set_font(run, face=R, size=11, color=GREY, bold=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    rPr.set("lang", "zh-CN")          # enables Chinese line-breaking rules (no punctuation at line start)
    rPr.set("altLang", "en-US")
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", face)


def runs_of(p, hl):
    """'abc **x** def' -> [(abc,{}),(x,hl),(def,{})]; lists of (text, opts) pass through."""
    if not isinstance(p, str):
        return p
    parts = re.split(r"\*\*(.+?)\*\*", p)
    return [(t, hl if i % 2 else {}) for i, t in enumerate(parts) if t]


# ------------------------------------------------------------------------------------- effects
def _spPr(shape):
    return shape._element.spPr


def _drop_fill(spPr):
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill"):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)


def _insert_fill(spPr, el):
    geom = spPr.find(qn("a:prstGeom"))
    if geom is None:
        geom = spPr.find(qn("a:custGeom"))
    (geom.addnext(el) if geom is not None else spPr.insert(0, el))


def _stops(stops):
    return "".join(f'<a:gs pos="{int(p * 1000)}"><a:srgbClr val="{c}"><a:alpha val="{int(a * 1000)}"/></a:srgbClr></a:gs>'
                   for p, c, a in stops)


def grad(shape, stops, angle=90):
    """Linear gradient. stops [(pos 0-100, hex, opacity 0-100)], angle 0 = left->right, 90 = top->bottom."""
    spPr = _spPr(shape)
    _drop_fill(spPr)
    _insert_fill(spPr, etree.fromstring(
        f'<a:gradFill xmlns:a="{A_NS}" rotWithShape="1"><a:gsLst>{_stops(stops)}</a:gsLst>'
        f'<a:lin ang="{int(angle * 60000)}" scaled="0"/></a:gradFill>'))
    return shape


def radial(shape, stops):
    spPr = _spPr(shape)
    _drop_fill(spPr)
    _insert_fill(spPr, etree.fromstring(
        f'<a:gradFill xmlns:a="{A_NS}" rotWithShape="1"><a:gsLst>{_stops(stops)}</a:gsLst>'
        f'<a:path path="circle"><a:fillToRect l="50000" t="50000" r="50000" b="50000"/></a:path></a:gradFill>'))
    return shape


def alpha(shape, pct):
    clr = _spPr(shape).find(qn("a:solidFill")).find(qn("a:srgbClr"))
    for a in clr.findall(qn("a:alpha")):
        clr.remove(a)
    etree.SubElement(clr, qn("a:alpha")).set("val", str(int(pct * 1000)))
    return shape


def shadow(shape, color=NAVY, opacity=22, blur=8, dist=2.5, angle=90):
    spPr = _spPr(shape)
    for e in spPr.findall(qn("a:effectLst")):
        spPr.remove(e)
    el = etree.fromstring(
        f'<a:effectLst xmlns:a="{A_NS}"><a:outerShdw blurRad="{int(blur * 12700)}" dist="{int(dist * 12700)}" '
        f'dir="{int(angle * 60000)}" algn="t" rotWithShape="0"><a:srgbClr val="{color}"><a:alpha val="{opacity * 1000}"/>'
        f'</a:srgbClr></a:outerShdw></a:effectLst>')
    ln = spPr.find(qn("a:ln"))
    (ln.addnext(el) if ln is not None else spPr.append(el))
    return shape


# ------------------------------------------------------------------------------------- deck
class Deck:
    def __init__(self, meta=None, kit=KIT):
        self.meta = meta or {}
        self.prs = Presentation(kit)
        self.hdr, self.icons, self.photos = {}, {}, {}
        for sl in self.prs.slides:
            for sh in sl.shapes:
                kind, _, key = sh.name.partition(":")
                if kind == "hdr":
                    self.hdr[key] = (sh._element, sl.part)
                elif kind == "icon":
                    self.icons[key] = (sh._element, sl.part)
                elif kind == "photo":
                    self.photos[key] = sh.image.blob
        lst = self.prs.slides._sldIdLst
        for sid in list(lst):
            self.prs.part.drop_rel(sid.rId)
            lst.remove(sid)
        self.layout = self.prs.slide_layouts[0]
        self._recolored = {}
        self.s = None

    # ---------------------------------------------------------------- frame
    def new(self, title="", section="", header=True, notes=None):
        """New slide with the template frame: hand icon + title + section tag + title rule, footer rule,
        disclaimer and signature. The frame is identical on every slide."""
        s = self.prs.slides.add_slide(self.layout)
        self.s = s
        if header:
            self._clone("title", title)
            self._clone("icon")
            self._clone("tag", section)
            self._clone("rule")
        self._clone("foot_rule")
        self._clone("disc", self.meta.get("disclaimer", "本报告所含的内容仅作参考。"))
        self._clone("sign", self.meta.get("sign", self.meta.get("org", "")))
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def _clone(self, key, text=None):
        el, part = self.hdr[key]
        el = copy.deepcopy(el)
        for node in el.iter():
            for attr in (qn("r:embed"), qn("r:link")):
                rid = node.get(attr)
                if rid:
                    node.set(attr, self.s.part.relate_to(part.related_part(rid), RT.IMAGE))
        if text is not None:
            runs = el.findall(".//" + qn("a:r"))
            runs[0].find(qn("a:t")).text = text
            for r in runs[1:]:
                r.getparent().remove(r)
        el.find(".//" + qn("p:cNvPr")).set("name", "frame:" + key)
        self.s.shapes._spTree.append(el)
        self._renumber()
        return el

    def _renumber(self):
        for n, c in enumerate(self.s.shapes._spTree.iter(qn("p:cNvPr")), start=2):
            c.set("id", str(n))

    # ---------------------------------------------------------------- primitives
    def rect(self, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE, dash=None, adj=None):
        sh = self.s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if adj is not None:
            for i, v in enumerate(adj if isinstance(adj, (list, tuple)) else [adj]):
                sh.adjustments[i] = v
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = rgb(line)
            sh.line.width = Pt(lw)
            if dash:
                sh.line.dash_style = dash
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        return sh

    def text(self, x, y, w, h, paras, size=FS_BODY, color=GREY, face=R, align="l", anchor="t",
             bold=False, spacing=None, after=0, wrap=True, shape=None, hl=None, vert=False):
        """paras: str | [str | [(text, opts)]]. '**x**' inside a str is highlighted with `hl`
        (default: OPPOSans B in brand red)."""
        hl = hl or {"face": B, "color": RED}
        sh = shape if shape is not None else self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = sh.text_frame
        tf.word_wrap = wrap
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        if vert:
            tf._txBody.find(qn("a:bodyPr")).set("vert", "eaVert")
        if isinstance(paras, str) or (paras and isinstance(paras[0], tuple)):
            paras = [paras]
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            if spacing:
                para.line_spacing = spacing
            if after and i < len(paras) - 1:
                para.space_after = Pt(after)
            for t, o in runs_of(p, hl):
                r = para.add_run()
                r.text = t
                set_font(r, o.get("face", face), o.get("size", size), o.get("color", color), o.get("bold", bold))
        return sh

    def box_text(self, x, y, w, h, txt, fill, color=WHITE, size=14, face=M, line=None, shape=MSO_SHAPE.RECTANGLE,
                 align="c", adj=None, lw=0.75, dash=None, hl=None):
        sh = self.rect(x, y, w, h, fill=fill, line=line, shape=shape, adj=adj, lw=lw, dash=dash)
        self.text(0, 0, 0, 0, txt, size=size, color=color, face=face, align=align, anchor="m", shape=sh, hl=hl)
        return sh

    def line(self, x1, y1, x2, y2, color=NAVY_B, w=0.75, dash=None, arrow=False):
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = rgb(color)
        c.line.width = Pt(w)
        if dash:
            c.line.dash_style = dash
        if arrow:
            te = etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd"))
            te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")
        return c

    def poly(self, pts, fill, stops=None, angle=90, line=None):
        fb = self.s.shapes.build_freeform(int(pts[0][0] * EMU), int(pts[0][1] * EMU), scale=1.0)
        fb.add_line_segments([(int(x * EMU), int(y * EMU)) for x, y in pts[1:]], close=True)
        sh = fb.convert_to_shape()
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
        if stops:
            grad(sh, stops, angle=angle)
        if line:
            sh.line.color.rgb = rgb(line); sh.line.width = Pt(0.75)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def chart(self, kind, data, x, y, w, h):
        return self.s.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data).chart

    # ---------------------------------------------------------------- template icons & photos
    def icon(self, key, x, y, size, color=None):
        """Template icon; `color` recolours the svg (and png fallback) to one flat colour."""
        if key not in self.icons:
            raise KeyError(f"unknown icon '{key}'. Available: {', '.join(sorted(self.icons))}")
        el, part = self.icons[key]
        el = copy.deepcopy(el)
        for node in el.iter():
            for attr in (qn("r:embed"), qn("r:link")):
                rid = node.get(attr)
                if rid:
                    tgt = part.related_part(rid)
                    if color:
                        tgt = self._recolor(tgt, color)
                    node.set(attr, self.s.part.relate_to(tgt, RT.IMAGE))
        xfrm = el.find(".//" + qn("a:xfrm"))
        xfrm.find(qn("a:off")).set("x", str(int(Inches(x)))); xfrm.find(qn("a:off")).set("y", str(int(Inches(y))))
        xfrm.find(qn("a:ext")).set("cx", str(int(Inches(size)))); xfrm.find(qn("a:ext")).set("cy", str(int(Inches(size))))
        self.s.shapes._spTree.append(el)
        self._renumber()
        return el

    def _recolor(self, part, color):
        key = (str(part.partname), color)
        if key in self._recolored:
            return self._recolored[key]
        ct = part.content_type
        if ct == "image/svg+xml":
            s = part.blob.decode("utf8")
            s = re.sub(r"#[0-9A-Fa-f]{6}\b|#[0-9A-Fa-f]{3}\b", "#" + color, s)
            s = s.replace("<svg ", f'<svg fill="#{color}" ', 1)
            # PowerPoint re-applies theme / "responsive" colours from MsftOfc* classes: drop them all
            s = re.sub(r'\sclass="MsftOfc[^"]*"', "", s)
            # some template icons are drawn semi-transparent (e.g. fill-opacity 0.2): a recoloured icon is always solid
            s = re.sub(r'(fill-opacity|stroke-opacity|opacity)="[0-9.]+"', r'\1="1"', s)
            s = re.sub(r"(fill-opacity|stroke-opacity|opacity)\s*:\s*[0-9.]+", r"\1:1", s)
            blob, ext = s.encode("utf8"), "svg"
        else:
            im = Image.open(io.BytesIO(part.blob)).convert("RGBA")
            solid = Image.new("RGBA", im.size, tuple(int(color[i:i + 2], 16) for i in (0, 2, 4)) + (255,))
            solid.putalpha(im.getchannel("A"))
            buf = io.BytesIO(); solid.save(buf, "PNG")
            blob, ext, ct = buf.getvalue(), "png", "image/png"
        pkg = self.prs.part.package
        new = Part(PackURI(pkg.next_partname("/ppt/media/image%d." + ext)), ct, pkg, blob)
        self._recolored[key] = new
        return new

    def photo(self, key, x, y, w, h, opacity=100, feather=0, max_px=1400):
        """Template photo, centre-cropped to the box; optionally faded and/or with feathered edges."""
        im = Image.open(io.BytesIO(self.photos[key]))
        has_alpha = im.mode in ("RGBA", "LA", "P")
        if feather or opacity < 100:
            max_px = min(max_px, 1000)       # faded backgrounds never need full resolution
        im.thumbnail((max_px, max_px))
        if feather:
            im = im.convert("RGBA")
            iw, ih = im.size
            f = int(min(iw, ih) * feather)
            mask = Image.new("L", im.size, 0)
            ImageDraw.Draw(mask).rectangle([f, f, iw - f, ih - f], fill=255)
            im.putalpha(ImageChops.multiply(im.getchannel("A"), mask.filter(ImageFilter.GaussianBlur(f / 2))))
            has_alpha = True
        buf = io.BytesIO()
        (im.convert("RGBA").save(buf, "PNG", optimize=True) if has_alpha else im.convert("RGB").save(buf, "JPEG", quality=85))
        buf.seek(0)
        pic = self.s.shapes.add_picture(buf, Inches(x), Inches(y), Inches(w), Inches(h))
        iw, ih = im.size
        box, img = w / h, iw / ih
        if img > box:
            pic.crop_left = pic.crop_right = (1 - box / img) / 2
        elif img < box:
            pic.crop_top = pic.crop_bottom = (1 - img / box) / 2
        if opacity < 100:
            etree.SubElement(pic._element.find(".//" + qn("a:blip")), qn("a:alphaModFix")).set("amt", str(opacity * 1000))
        return pic

    def photo_fit(self, key, x, y, w):
        iw, ih = Image.open(io.BytesIO(self.photos[key])).size
        return self.photo(key, x, y, w, w * ih / iw)

    def save(self, path):
        self.prs.save(path)
